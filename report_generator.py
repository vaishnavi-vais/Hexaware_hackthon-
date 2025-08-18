import pandas as pd
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
import seaborn as sns
import io
import base64
from datetime import datetime
import os
from typing import Dict, List, Any

class ReportGenerator:
    def __init__(self):
        self.styles = getSampleStyleSheet()
        self.custom_styles = self._create_custom_styles()
        
    def _create_custom_styles(self):
        """Create custom styles for reports"""
        styles = {}
        styles['CustomTitle'] = ParagraphStyle(
            'CustomTitle',
            parent=self.styles['Title'],
            fontSize=24,
            spaceAfter=30,
            textColor=colors.HexColor('#2E86AB')
        )
        styles['CustomHeading'] = ParagraphStyle(
            'CustomHeading',
            parent=self.styles['Heading1'],
            fontSize=16,
            spaceAfter=12,
            textColor=colors.HexColor('#A23B72')
        )
        return styles
    
    def generate_question_bank_excel(self, questions: List[Dict], filename: str) -> str:
        """Generate Excel file with question bank"""
        df_data = []
        
        for i, question in enumerate(questions, 1):
            if question['type'] == 'mcq':
                df_data.append({
                    'Question_ID': f"Q{i:03d}",
                    'Type': 'MCQ',
                    'Question': question['question'],
                    'Option_A': question['options'][0] if len(question['options']) > 0 else '',
                    'Option_B': question['options'][1] if len(question['options']) > 1 else '',
                    'Option_C': question['options'][2] if len(question['options']) > 2 else '',
                    'Option_D': question['options'][3] if len(question['options']) > 3 else '',
                    'Correct_Answer': chr(65 + question['correct_answer']),  # A, B, C, D
                    'Difficulty': question['difficulty'].title(),
                    'Topic': question['topic'],
                    'Explanation': question.get('explanation', '')
                })
            elif question['type'] == 'coding':
                df_data.append({
                    'Question_ID': f"Q{i:03d}",
                    'Type': 'Coding',
                    'Question': question['question'],
                    'Language': question['language'],
                    'Sample_Code': question['sample_code'],
                    'Difficulty': question['difficulty'].title(),
                    'Topic': question['topic'],
                    'Test_Cases': str(question.get('test_cases', []))
                })
            elif question['type'] == 'case_study':
                df_data.append({
                    'Question_ID': f"Q{i:03d}",
                    'Type': 'Case Study',
                    'Question': question['question'],
                    'Scenario': question['scenario'],
                    'Domain': question['domain'],
                    'Difficulty': question['difficulty'].title(),
                    'Topic': question['topic'],
                    'Evaluation_Criteria': ', '.join(question.get('evaluation_criteria', []))
                })
        
        df = pd.DataFrame(df_data)
        
        # Create Excel file with formatting
        with pd.ExcelWriter(filename, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Question Bank', index=False)
            
            # Get the workbook and worksheet
            workbook = writer.book
            worksheet = writer.sheets['Question Bank']
            
            # Auto-adjust column widths
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        return filename
    
    def generate_feedback_report_pdf(self, feedback_data: List[Dict], filename: str) -> str:
        """Generate PDF feedback analysis report"""
        doc = SimpleDocTemplate(filename, pagesize=A4)
        story = []
        
        # Title
        title = Paragraph("Feedback Analysis Report", self.custom_styles['CustomTitle'])
        story.append(title)
        story.append(Spacer(1, 20))
        
        # Report metadata
        metadata = [
            ['Report Generated:', datetime.now().strftime('%Y-%m-%d %H:%M:%S')],
            ['Total Feedback Items:', str(len(feedback_data))],
            ['Analysis Period:', 'Last 30 days']
        ]
        
        metadata_table = Table(metadata, colWidths=[2*inch, 3*inch])
        metadata_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ]))
        
        story.append(metadata_table)
        story.append(Spacer(1, 20))
        
        # Sentiment Analysis Summary
        sentiment_counts = {'positive': 0, 'negative': 0, 'neutral': 0}
        for feedback in feedback_data:
            sentiment = feedback.get('sentiment', 'neutral')
            sentiment_counts[sentiment] += 1
        
        story.append(Paragraph("Sentiment Analysis Summary", self.custom_styles['CustomHeading']))
        
        sentiment_data = [
            ['Sentiment', 'Count', 'Percentage'],
            ['Positive', str(sentiment_counts['positive']), f"{sentiment_counts['positive']/len(feedback_data)*100:.1f}%"],
            ['Negative', str(sentiment_counts['negative']), f"{sentiment_counts['negative']/len(feedback_data)*100:.1f}%"],
            ['Neutral', str(sentiment_counts['neutral']), f"{sentiment_counts['neutral']/len(feedback_data)*100:.1f}%"]
        ]
        
        sentiment_table = Table(sentiment_data, colWidths=[2*inch, 1*inch, 1.5*inch])
        sentiment_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        story.append(sentiment_table)
        story.append(Spacer(1, 20))
        
        # Key Insights
        story.append(Paragraph("Key Insights", self.custom_styles['CustomHeading']))
        
        all_insights = []
        for feedback in feedback_data:
            all_insights.extend(feedback.get('insights', []))
        
        # Get top 5 most common insights
        insight_counts = {}
        for insight in all_insights:
            insight_counts[insight] = insight_counts.get(insight, 0) + 1
        
        top_insights = sorted(insight_counts.items(), key=lambda x: x[1], reverse=True)[:5]
        
        for i, (insight, count) in enumerate(top_insights, 1):
            story.append(Paragraph(f"{i}. {insight} (mentioned {count} times)", self.styles['Normal']))
            story.append(Spacer(1, 6))
        
        story.append(Spacer(1, 20))
        
        # Recommendations
        story.append(Paragraph("Recommendations", self.custom_styles['CustomHeading']))
        
        recommendations = [
            "Increase focus on areas with negative feedback",
            "Leverage positive feedback patterns for improvement",
            "Implement regular feedback collection cycles",
            "Create targeted improvement plans based on insights"
        ]
        
        for i, rec in enumerate(recommendations, 1):
            story.append(Paragraph(f"{i}. {rec}", self.styles['Normal']))
            story.append(Spacer(1, 6))
        
        doc.build(story)
        return filename
    
    def generate_curriculum_mapping_docx(self, mapping_data: Dict, filename: str) -> str:
        """Generate DOCX curriculum mapping document"""
        doc = Document()
        
        # Title
        title = doc.add_heading('Curriculum Mapping Report', 0)
        title.alignment = 1  # Center alignment
        
        # Metadata
        doc.add_heading('Report Information', level=1)
        metadata_table = doc.add_table(rows=3, cols=2)
        metadata_table.style = 'Table Grid'
        
        metadata_table.cell(0, 0).text = 'Generated On'
        metadata_table.cell(0, 1).text = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        metadata_table.cell(1, 0).text = 'Skill Categories'
        metadata_table.cell(1, 1).text = str(len(mapping_data))
        metadata_table.cell(2, 0).text = 'Analysis Type'
        metadata_table.cell(2, 1).text = 'AI-Powered Curriculum Mapping'
        
        doc.add_page_break()
        
        # Skill Categories
        doc.add_heading('Skill Categories and Mapping', level=1)
        
        for category, details in mapping_data.items():
            doc.add_heading(category.replace('_', ' ').title(), level=2)
            
            # Keywords
            doc.add_heading('Key Topics:', level=3)
            keywords_para = doc.add_paragraph()
            keywords_para.add_run(', '.join(details['keywords']))
            
            # Difficulty
            doc.add_heading('Estimated Difficulty:', level=3)
            difficulty_para = doc.add_paragraph()
            difficulty_para.add_run(details['difficulty_estimate'].title())
            
            # Content Items
            doc.add_heading('Related Content Items:', level=3)
            content_para = doc.add_paragraph()
            content_para.add_run(f"{len(details['content_items'])} curriculum items mapped to this category")
            
            doc.add_paragraph()  # Add spacing
        
        # Summary
        doc.add_page_break()
        doc.add_heading('Summary and Recommendations', level=1)
        
        summary_points = [
            f"Total of {len(mapping_data)} skill categories identified",
            "Curriculum content successfully mapped to competency areas",
            "Difficulty levels estimated using AI analysis",
            "Ready for question generation and assessment planning"
        ]
        
        for point in summary_points:
            doc.add_paragraph(point, style='List Bullet')
        
        doc.save(filename)
        return filename
    
    def generate_learning_plan_docx(self, user_data: Dict, plan_data: Dict, filename: str) -> str:
        """Generate personalized learning plan in DOCX format"""
        doc = Document()
        
        # Title
        title = doc.add_heading('Personalized Learning Plan', 0)
        title.alignment = 1
        
        # User Information
        doc.add_heading('Learner Information', level=1)
        user_table = doc.add_table(rows=4, cols=2)
        user_table.style = 'Table Grid'
        
        user_table.cell(0, 0).text = 'Name'
        user_table.cell(0, 1).text = user_data.get('username', 'N/A')
        user_table.cell(1, 0).text = 'Department'
        user_table.cell(1, 1).text = user_data.get('department', 'N/A')
        user_table.cell(2, 0).text = 'Role'
        user_table.cell(2, 1).text = user_data.get('role', 'N/A')
        user_table.cell(3, 0).text = 'Plan Generated'
        user_table.cell(3, 1).text = datetime.now().strftime('%Y-%m-%d')
        
        doc.add_page_break()
        
        # Learning Objectives
        doc.add_heading('Learning Objectives', level=1)
        objectives = plan_data.get('objectives', [
            'Master key concepts in assigned curriculum',
            'Develop practical skills through hands-on exercises',
            'Achieve proficiency in assessment areas',
            'Apply knowledge to real-world scenarios'
        ])
        
        for objective in objectives:
            doc.add_paragraph(objective, style='List Bullet')
        
        # Learning Path
        doc.add_heading('Recommended Learning Path', level=1)
        
        learning_modules = plan_data.get('modules', [
            {'name': 'Foundation Concepts', 'duration': '2 weeks', 'difficulty': 'Easy'},
            {'name': 'Intermediate Applications', 'duration': '3 weeks', 'difficulty': 'Medium'},
            {'name': 'Advanced Topics', 'duration': '2 weeks', 'difficulty': 'Hard'},
            {'name': 'Practical Projects', 'duration': '1 week', 'difficulty': 'Medium'}
        ])
        
        for i, module in enumerate(learning_modules, 1):
            doc.add_heading(f"Module {i}: {module['name']}", level=2)
            
            module_table = doc.add_table(rows=2, cols=2)
            module_table.style = 'Table Grid'
            
            module_table.cell(0, 0).text = 'Duration'
            module_table.cell(0, 1).text = module['duration']
            module_table.cell(1, 0).text = 'Difficulty'
            module_table.cell(1, 1).text = module['difficulty']
            
            doc.add_paragraph()
        
        # Assessment Schedule
        doc.add_heading('Assessment Schedule', level=1)
        
        assessments = plan_data.get('assessments', [
            {'name': 'Module 1 Quiz', 'type': 'MCQ', 'week': 2},
            {'name': 'Coding Challenge', 'type': 'Practical', 'week': 4},
            {'name': 'Case Study Analysis', 'type': 'Written', 'week': 6},
            {'name': 'Final Assessment', 'type': 'Comprehensive', 'week': 8}
        ])
        
        assessment_table = doc.add_table(rows=len(assessments) + 1, cols=3)
        assessment_table.style = 'Table Grid'
        
        # Headers
        assessment_table.cell(0, 0).text = 'Assessment'
        assessment_table.cell(0, 1).text = 'Type'
        assessment_table.cell(0, 2).text = 'Week'
        
        # Data
        for i, assessment in enumerate(assessments, 1):
            assessment_table.cell(i, 0).text = assessment['name']
            assessment_table.cell(i, 1).text = assessment['type']
            assessment_table.cell(i, 2).text = str(assessment['week'])
        
        doc.save(filename)
        return filename
    
    def generate_assessment_completion_report_pdf(self, assessment_data: List[Dict], filename: str) -> str:
        """Generate assessment completion report in PDF format"""
        doc = SimpleDocTemplate(filename, pagesize=A4)
        story = []
        
        # Title
        title = Paragraph("Assessment Completion Report", self.custom_styles['CustomTitle'])
        story.append(title)
        story.append(Spacer(1, 20))
        
        # Summary Statistics
        total_assessments = len(assessment_data)
        completed_assessments = len([a for a in assessment_data if a.get('completed_at')])
        avg_score = sum([a.get('score', 0) for a in assessment_data if a.get('score')]) / max(1, len([a for a in assessment_data if a.get('score')]))
        
        summary_data = [
            ['Metric', 'Value'],
            ['Total Assessments', str(total_assessments)],
            ['Completed Assessments', str(completed_assessments)],
            ['Completion Rate', f"{completed_assessments/max(1, total_assessments)*100:.1f}%"],
            ['Average Score', f"{avg_score:.1f}%"]
        ]
        
        summary_table = Table(summary_data, colWidths=[3*inch, 2*inch])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        story.append(summary_table)
        story.append(Spacer(1, 30))
        
        # Detailed Results
        story.append(Paragraph("Detailed Assessment Results", self.custom_styles['CustomHeading']))
        
        if assessment_data:
            # Create detailed table
            detailed_data = [['User', 'Score', 'Questions', 'Correct', 'Time (min)', 'Status']]
            
            for assessment in assessment_data[:20]:  # Limit to first 20 for space
                status = 'Completed' if assessment.get('completed_at') else 'In Progress'
                detailed_data.append([
                    assessment.get('username', 'Unknown'),
                    f"{assessment.get('score', 0):.1f}%",
                    str(assessment.get('total_questions', 0)),
                    str(assessment.get('correct_answers', 0)),
                    str(assessment.get('time_taken', 0)),
                    status
                ])
            
            detailed_table = Table(detailed_data, colWidths=[1.2*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch, 1*inch])
            detailed_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(detailed_table)
        
        doc.build(story)
        return filename
